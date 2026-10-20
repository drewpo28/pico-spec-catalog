"""ZX Spectrum Next Database adapter — free Next games and demos on itch.io.

Tree exposed to the device:

    Games / <letter> / <entry>              letters: 0-9, A..Z
    Games / <letter> / <title> / <entry>    a game with several runnable variants
    Demos / ...                             same shape

The list comes from Leonis' "ZX Spectrum Next Database"
(https://leonisx.itch.io/zx-spectrum-next-db), an .xlsx published on itch.io
and refreshed by hand. Everything below was checked from a GitHub runner on
2026-10-18 (itch.io is not reachable from the dev container):

  - the page is a "name your price" book; `POST <page>/download_url` returns a
    keyed download page, whose upload has a `data-upload_id`. The file itself:
    `POST <page>/file/<upload_id>?source=game_download` (empty body — no csrf,
    no cookie, no key needed) answers JSON {"url": <signed Cloudflare R2 URL>},
    and that URL is valid for 60 SECONDS (X-Amz-Expires=60). This holds for the
    game uploads too, so a locator here can never be a plain URL: it names the
    page and the upload, and the device runs the POST itself right before it
    downloads (see the firmware's HttpCatalogFs, "itch:" locators).
  - sheets "Games" and "Demos": header in row 2, data from row 3. Column
    "£Digital" holds the price ("Free", "NYOP", "£8.95", "Free (was $5.99..."),
    and the cell's HYPERLINK is the release page. Only Free*/NYOP rows whose link
    is an itch.io game page are taken (~215), and the page's own
    init_GameDownload JSON decides: min_price must be 0 (two "NYOP" games have a
    minimum, the sheet does not say so).
  - a game page lists its uploads with data-upload_id, except "Download Now"
    (NYOP) pages, which list them without ids — those come from the keyed
    download page, the same way as the database itself.
  - uploads: mostly .zip (a game is usually MANY files — a .bas or .nex plus a
    data/ folder of .nxi/.spr/.nxt, loaded by relative path), then .nex, .tap,
    names without an extension (the real name is in the R2 Content-Disposition),
    pdf/txt manuals, other platforms' builds, sources, the odd .rar/.7z.

The database misses many itch-only releases, so the free listings of the Next
tags (TAGS; checked 2026-10-20: ~187 games, ~79 not in the database, ~57 of
them runnable) are merged in. `<tag list>?page=N&format=json` answers
{"page", "num_items", "content": <grid html>}, 36 cards a page. A card's
attribute ORDER changes between responses, so cards are cut at their
`game_cell` div and read by `data-label="game:<id>:title"` anchors, not by
attribute position. The database wins for a page it lists (title, developer,
sheet); a tag-only page goes to Demos when it carries a demo tag, else Games,
and passes the same free check and upload pipeline — the noisy "next" tag
needs no filter of its own, a page with nothing runnable lists nothing. A
failed tag listing is only a warning; MIN_GAMES counts database rows alone.

Per upload the adapter decides what the device should RUN:
  - .nex / .snx / .bas uploaded on their own: that file;
  - .zip: its central directory is read with HTTP Range requests (no full
    download — some archives are 30-50 MB) and every runnable member at the
    shallowest level (.nex, else .bas, else .tap) becomes one entry; e.g.
    Aliens: Neoplasma ships EN/RU/SP .nex, so it lists three. A zip with nothing
    runnable (sources, browser builds, nested zips) is dropped.
The member to run is carried as "?run=<member>" after the file name; the device
unpacks the whole archive into /tmp/next/<game>/ and launches that member from
there through NextZXOS, so relative data paths work.

Files are NOT mirrored: listings only, every download goes to itch.io.

Knobs (env):
  NEXTDB_REQ_GAP  minimum seconds between requests to itch.io (default 0.4)
  NEXTDB_CACHE    JSON cache of per-upload results (zip members, real names),
                  keyed by upload id + size, so a nightly run only reads new
                  or changed uploads (default .cache/nextdb.json; "" = off)
"""

from __future__ import annotations

import html
import io
import json
import os
import re
import struct
import time
import zipfile
from urllib.parse import quote, urlparse

from .base import Adapter, Entry, SourceDown, http_client

DB_PAGE = "https://leonisx.itch.io/zx-spectrum-next-db"
SHEETS = ("Games", "Demos")
TAGS = ("zx-spectrum-next", "spectrum-next", "zx-next", "zxnext", "next")
TAG_LIST = "https://itch.io/games/free/tag-{}"
TAG_PAGE_SIZE = 36
TAG_MAX_PAGES = 20

UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/120 Safari/537.36"

GAME_PAGE = re.compile(r"^https?://([a-z0-9-]+\.itch\.io)/([A-Za-z0-9_-]+)/?$", re.I)
CSRF = re.compile(r'<meta name="csrf_token" value="([^"]+)"')
MIN_PRICE = re.compile(r'"min_price":(\d+)')
UPLOAD_BLK = re.compile(r'<div class="upload">(.*?)</div>\s*</div>', re.S)
UPLOAD_ID = re.compile(r'data-upload_id="(\d+)"')
UPLOAD_NAME = re.compile(r'class="name"[^>]*title="([^"]*)"')
UPLOAD_NAME2 = re.compile(r'class="name"[^>]*>([^<]*)<')
UPLOAD_SIZE = re.compile(r'class="file_size"[^>]*><span>([^<]*)<')
PAGE_TITLE = re.compile(r'<meta property="og:title" content="([^"]*)"|<title>([^<]*)</title>')
DISPO = re.compile(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)"?', re.I)

# Tag listings: one card per `game_cell` div, attributes in any order.
CARD_SPLIT = re.compile(r'<div[^>]*class="game_cell ')
CARD_TITLE = re.compile(r'<a [^>]*data-label="game:\d+:title"[^>]*>([^<]*)<')
CARD_AUTHOR = re.compile(r'class="game_author"[^>]*>\s*<a [^>]*>([^<]*)<')
HREF = re.compile(r'href="([^"]+)"')
PAGE_TAGS = re.compile(r'href="https://itch\.io/games/tag-([a-z0-9-]+)"')
DEMO_TAGS = {"demo", "demoscene", "demoscene-production", "intro"}
# "Humbug - ZX Next", "Infestation (ZX Spectrum Next)", "... - Spectrum Next".
# Only after a separator: "Scramble Next" keeps its name.
NEXT_SUFFIX = re.compile(
    r"\s*(?:\(|-|–|—)\s*(?:for\s+)?(?:the\s+)?(?:zx\s+spectrum|zx|spectrum)\s+next"
    r"(?:\s+edition)?\s*\)?\s*$", re.I)

RUN_EXT = ("nex", "snx", "bas")
TAPE_EXT = ("tap", "tzx")
KEEP_EXT = RUN_EXT + TAPE_EXT + ("zip",)

# An upload whose NAME says it is not the Next build, or not a program at all.
NOT_PROGRAM = re.compile(
    r"\b(source|src|code|manual|guide|instructions?|readme|cheat|hints?|"
    r"solution|map|tips|poster|cover|artwork|soundtrack|ost|feelies)\b", re.I)
OTHER_PLATFORM = re.compile(
    r"\b(msx2?|c64|commodore|amiga|atari|dos|ql|sqlux|smsq|mac|macos|windows|win|"
    r"linux|pet|vic-?20|dragon|psp|sam|coup[eé]|bbc|cpc|amstrad|plus4|zx\s*touch|"
    r"48k?|128k?|48/128|esxdos|dsk|zesarux|cgleches)\b", re.I)
NEXT_WORD = re.compile(r"next", re.I)

# Members of an archive that are never the program even if they end in .bas.
SKIP_DIR = re.compile(r"(^|/)(__macosx|src|source|sources|tools?|procbanks|"
                      r"zesarux|emulator|godot|res)(/|$)", re.I)

MIN_GAMES = 100      # fewer than this means the sheet moved, not that games vanished
MAX_VARIANTS = 6     # runnable members listed per archive


def _text(s) -> str:
    s = html.unescape(str(s or ""))
    s = s.replace("/", "-").replace("\t", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()


def _letter(title: str) -> str:
    c = title[:1].upper()
    return c if "A" <= c <= "Z" else "0-9"


def _ext(name: str) -> str:
    base = name.rsplit("/", 1)[-1]
    return base.rsplit(".", 1)[-1].lower() if "." in base else ""


def _size(s: str) -> int:
    m = re.match(r"\s*([\d.,]+)\s*(bytes|kB|MB|GB)", s or "")
    if not m:
        return 0
    v = float(m.group(1).replace(",", ""))
    return int(v * {"bytes": 1, "kB": 1024, "MB": 1 << 20, "GB": 1 << 30}[m.group(2)])


def uploads_of(page: str) -> list[tuple[str | None, str, int]]:
    """(upload_id or None, display name, size) per upload block of an itch page."""
    out = []
    for m in UPLOAD_BLK.finditer(page):
        b = m.group(1)
        uid = UPLOAD_ID.search(b)
        nm = UPLOAD_NAME.search(b) or UPLOAD_NAME2.search(b)
        sz = UPLOAD_SIZE.search(b)
        out.append((uid.group(1) if uid else None,
                    html.unescape(nm.group(1)).strip() if nm else "",
                    _size(sz.group(1)) if sz else 0))
    return out


def tag_cards(grid: str) -> list[tuple[str, str, str]]:
    """(game page URL, title, author) per card of a tag listing's grid html."""
    out = []
    for cell in CARD_SPLIT.split(grid)[1:]:
        a = CARD_TITLE.search(cell)
        if not a:
            continue
        h = HREF.search(a.group(0))
        if not h or not GAME_PAGE.match(h.group(1)):
            continue
        au = CARD_AUTHOR.search(cell)
        title = NEXT_SUFFIX.sub("", _text(a.group(1))) or _text(a.group(1))
        out.append((h.group(1).rstrip("/"), title, _text(au.group(1)) if au else ""))
    return out


def pick_uploads(ups: list[tuple[str, str, int]]) -> list[tuple[str, str, int]]:
    """Drop manuals, sources and other platforms' builds; prefer the 'Next' ones.

    `ups` carries REAL file names (extension known). An upload is a candidate
    when its extension can run on the Next (or is a zip that may hold one) and
    neither its name nor its file name marks it as something else."""
    cand = []
    for uid, name, size in ups:
        ext = _ext(name)
        if ext not in KEEP_EXT:
            continue
        if NOT_PROGRAM.search(name.replace("_", " ")):
            continue
        cand.append((uid, name, size))
    nexty = [u for u in cand if NEXT_WORD.search(u[1])]
    if nexty:
        rest = [u for u in cand if u not in nexty and _ext(u[1]) in RUN_EXT]
        cand = nexty + rest
    else:
        cand = [u for u in cand if not OTHER_PLATFORM.search(u[1].replace("_", " "))]
    # A bare tape is a 48K/128K release unless it says otherwise; inside a zip a
    # .tap does run on the Next (DAAD adventures load their .nxt from beside it).
    return [u for u in cand if _ext(u[1]) not in TAPE_EXT or NEXT_WORD.search(u[1])]


def _member_name(i: "zipfile.ZipInfo") -> str:
    """A member's name as the device's unzip writes it: '/' separators, and a
    name without the UTF-8 flag read as CP866 (zipfile assumes CP437)."""
    n = i.filename
    if not i.flag_bits & 0x800:
        try:
            n = n.encode("cp437").decode("cp866")
        except UnicodeError:
            pass
    return n.replace("\\", "/")


def run_members(names: list[str]) -> list[str]:
    """The members of an archive worth launching, shallowest level only.

    .nex first (a NEX is the release), then .bas, then a .tap — one kind only,
    all candidates of that kind at the shallowest directory depth, autoexec.bas
    preferred when present."""
    files = [n for n in names if not n.endswith("/") and not SKIP_DIR.search(n)]
    for kind in (("nex", "snx"), ("bas",), ("tap",)):
        hits = [n for n in files if _ext(n) in kind]
        if not hits:
            continue
        depth = min(n.count("/") for n in hits)
        top = sorted(n for n in hits if n.count("/") == depth)
        auto = [n for n in top if n.rsplit("/", 1)[-1].lower() == "autoexec.bas"]
        return auto or top[:MAX_VARIANTS]
    return []


class RangeFile(io.RawIOBase):
    """A seekable read-only file over HTTP Range requests (for zipfile)."""

    def __init__(self, get_url, size: int):
        self._get_url = get_url    # callable → a fresh signed URL
        self._size = size
        self._pos = 0
        self._url = get_url()
        self._t = time.monotonic()

    def readable(self): return True
    def seekable(self): return True
    def tell(self): return self._pos

    def seek(self, off, whence=0):
        self._pos = {0: off, 1: self._pos + off, 2: self._size + off}[whence]
        return self._pos

    def read(self, n=-1):
        if n is None or n < 0:
            n = self._size - self._pos
        n = min(n, self._size - self._pos)
        if n <= 0:
            return b""
        if time.monotonic() - self._t > 45:          # the signature lasts 60 s
            self._url, self._t = self._get_url(), time.monotonic()
        r = self._get_url.client.get(self._url, headers={"Range": f"bytes={self._pos}-{self._pos + n - 1}"})
        r.raise_for_status()
        data = r.content[:n]
        self._pos += len(data)
        return data

    def readinto(self, b):
        d = self.read(len(b))
        b[:len(d)] = d
        return len(d)


class NextDbAdapter(Adapter):
    id = "nextdb"
    name = "ZX Spectrum Next (itch.io)"

    def __init__(self):
        self._client = http_client(headers={"User-Agent": UA}, timeout=60.0, follow_redirects=True)
        self._gap = float(os.environ.get("NEXTDB_REQ_GAP", "0.4"))
        self._last = 0.0
        self._cache_path = os.environ.get("NEXTDB_CACHE", ".cache/nextdb.json")
        self._cache: dict = {}
        if self._cache_path and os.path.exists(self._cache_path):
            try:
                with open(self._cache_path, encoding="utf-8") as fh:
                    self._cache = json.load(fh)
            except Exception:  # noqa: BLE001 — a broken cache only costs a re-read
                self._cache = {}
        self._tree: dict[str, list[Entry]] | None = None

    # ── HTTP ────────────────────────────────────────────────────────────────
    def _req(self, method: str, url: str, **kw):
        err = None
        for attempt in range(4):
            wait = self._last + self._gap - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            try:
                r = self._client.request(method, url, **kw)
                self._last = time.monotonic()
                if r.status_code in (429, 500, 502, 503, 504):
                    raise RuntimeError(f"HTTP {r.status_code}")
                return r
            except Exception as e:  # noqa: BLE001 — retried, then reported
                self._last = time.monotonic()
                err = e
                time.sleep(2 * (attempt + 1))
        raise RuntimeError(f"nextdb: {method} {url}: {err}")

    def _signed(self, page: str, uid: str) -> str:
        r = self._req("POST", f"{page}/file/{uid}", params={"source": "game_download"})
        try:
            return r.json()["url"]
        except Exception:  # noqa: BLE001
            raise RuntimeError(f"nextdb: no download url for {page} #{uid}: {r.text[:120]}")

    def _keyed_uploads(self, page: str, html_text: str) -> list[tuple[str | None, str, int]]:
        m = CSRF.search(html_text)
        r = self._req("POST", f"{page}/download_url", data={"csrf_token": m.group(1) if m else ""})
        try:
            url = r.json()["url"]
        except Exception:  # noqa: BLE001
            return []
        return uploads_of(self._req("GET", url).text)

    # ── the database ────────────────────────────────────────────────────────
    def _rows(self) -> list[tuple[str, str, str, str]]:
        """(sheet, title, developer, game page URL) for every free itch release."""
        try:
            page = self._req("GET", DB_PAGE).text
            ups = [u for u in self._keyed_uploads(DB_PAGE, page) if u[0] and u[1].lower().endswith(".xlsx")]
            if not ups:
                raise RuntimeError("no .xlsx upload on the database page")
            data = self._req("GET", self._signed(DB_PAGE, ups[0][0])).content
        except Exception as e:  # noqa: BLE001
            raise SourceDown(f"nextdb: database download failed: {e}")
        import openpyxl  # lazy: only this source needs it
        wb = openpyxl.load_workbook(io.BytesIO(data))
        out = []
        for sheet in SHEETS:
            if sheet not in wb.sheetnames:
                raise SourceDown(f"nextdb: sheet {sheet!r} missing — layout changed?")
            ws = wb[sheet]
            hdr = [str(c.value or "").strip() for c in ws[2]]
            try:
                tc, pc = hdr.index("Title"), hdr.index("£Digital")
            except ValueError:
                raise SourceDown(f"nextdb: {sheet}: header {hdr[:12]} has no Title/£Digital")
            dc = next((i for i, h in enumerate(hdr) if h.startswith("Developer")), None)
            for row in ws.iter_rows(min_row=3):
                cell = row[pc]
                price = str(cell.value or "").strip()
                link = (cell.hyperlink.target or "").strip() if cell.hyperlink else ""
                if not (price.startswith("Free") or price.startswith("NYOP")):
                    continue
                if not GAME_PAGE.match(link) or link.rstrip("/") == DB_PAGE:
                    continue
                title = _text(row[tc].value)
                if not title:
                    continue
                dev = _text(row[dc].value) if dc is not None else ""
                out.append((sheet, title, dev, link.rstrip("/")))
        return out

    def _tag_rows(self, known: set[str]) -> list[tuple[None, str, str, str]]:
        """(None, title, author, game page URL) for every free page under TAGS
        that `known` (lower-cased database pages) does not list. The sheet is
        decided once the page itself is read."""
        out, seen = [], set(known) | {DB_PAGE.lower()}
        for tag in TAGS:
            n = 0
            try:
                for p in range(1, TAG_MAX_PAGES + 1):
                    r = self._req("GET", TAG_LIST.format(tag), params={"page": p, "format": "json"})
                    j = r.json()
                    cards = tag_cards(j.get("content") or "")
                    for page, title, dev in cards:
                        if page.lower() in seen:
                            continue
                        seen.add(page.lower())
                        out.append((None, title, dev, page))
                        n += 1
                    if int(j.get("num_items") or 0) < TAG_PAGE_SIZE:
                        break
            except Exception as e:  # noqa: BLE001 — the tags only add to the database
                print(f"  ! nextdb: tag {tag}: {e}")
            print(f"  nextdb tag {tag}: {n} new pages")
        return out

    # ── one upload ──────────────────────────────────────────────────────────
    def _upload_info(self, page: str, uid: str, name: str, size: int) -> dict:
        """{"file": real name, "run": [members]} for one upload (cached)."""
        key = f"{uid}:{size}"
        hit = self._cache.get(key)
        if hit is not None:
            return hit
        info = {"file": name, "run": []}
        try:
            if _ext(name) not in KEEP_EXT + ("pdf", "txt", "png", "jpg", "rar", "7z", "exe", "dmg", "rtf"):
                # No usable extension in the display name: the real one is in the
                # R2 Content-Disposition. One byte is enough to see the headers.
                r = self._req("GET", self._signed(page, uid), headers={"Range": "bytes=0-0"})
                m = DISPO.search(r.headers.get("content-disposition", ""))
                if m:
                    info["file"] = m.group(1).strip()
            if _ext(info["file"]) == "zip":
                getter = lambda: self._signed(page, uid)  # noqa: E731
                getter.client = self._client
                head = self._req("GET", getter(), headers={"Range": "bytes=0-0"})
                total = int((head.headers.get("content-range", "").rsplit("/", 1) + ["0"])[1] or 0)
                if total <= 0:
                    raise RuntimeError("no Content-Range")
                with zipfile.ZipFile(RangeFile(getter, total)) as zf:
                    info["run"] = run_members([_member_name(i) for i in zf.infolist()])
            elif _ext(info["file"]) in RUN_EXT or _ext(info["file"]) in TAPE_EXT:
                info["run"] = [""]
        except Exception as e:  # noqa: BLE001 — one bad upload is not a dead site
            print(f"  ! nextdb: {page} #{uid} {name!r}: {e}")
            return info          # not cached: try again next night
        self._cache[key] = info
        return info

    def _save_cache(self) -> None:
        if not self._cache_path:
            return
        os.makedirs(os.path.dirname(self._cache_path) or ".", exist_ok=True)
        with open(self._cache_path, "w", encoding="utf-8") as fh:
            json.dump(self._cache, fh)

    # ── the tree ────────────────────────────────────────────────────────────
    def _load(self) -> dict[str, list[Entry]]:
        if self._tree is not None:
            return self._tree
        rows = self._rows()
        rows += self._tag_rows({r[3].lower() for r in rows})
        per_page: dict[tuple[str, str], int] = {}
        for sheet, _, _, page in rows:
            per_page[(sheet, page)] = per_page.get((sheet, page), 0) + 1
        tree: dict[str, list[Entry]] = {}
        seen_page: set[tuple[str, str]] = set()
        games = tagged = 0
        for sheet, title, dev, page in rows:
            if (sheet, page) in seen_page:      # one itch page, several sheet rows
                continue                        # (Nexter issues list each game)
            seen_page.add((sheet, page))
            try:
                html_text = self._req("GET", page).text
            except Exception as e:  # noqa: BLE001
                print(f"  ! nextdb: {page}: {e}")
                continue
            if per_page[(sheet, page)] > 1:
                # Several sheet rows share one release page (a compilation such as
                # Sunteam's "Nexter" issues): name it after the page, not one game.
                tm = PAGE_TITLE.search(html_text)
                if tm:
                    t = _text(tm.group(1) or tm.group(2))
                    t = re.sub(r"\s+by\s+[^-]+$", "", t).strip()
                    title, dev = t or title, ""
            m = MIN_PRICE.search(html_text)
            if not m or int(m.group(1)) != 0:
                continue                        # not free after all
            from_db = sheet is not None
            if not from_db:                     # a tag-only page
                sheet = "Demos" if DEMO_TAGS & set(PAGE_TAGS.findall(html_text)) else "Games"
            ups = uploads_of(html_text)
            if ups and not all(u[0] for u in ups):
                ups = self._keyed_uploads(page, html_text) or ups
            ups = [u for u in ups if u[0]]
            real = []
            for uid, name, size in ups:
                info = self._upload_info(page, uid, name, size)
                real.append((uid, info["file"], size, info, name))
            chosen = {u[0] for u in pick_uploads([(a, b, c) for a, b, c, _, _ in real])}
            items = []
            for uid, fname, size, info, shown in real:
                if uid not in chosen:
                    continue
                for member in info["run"]:
                    items.append((uid, fname, size, member, shown, len(info["run"])))
            if not items:
                continue
            host, slug = GAME_PAGE.match(page).groups()
            base = f"{sheet}/{_letter(title)}"
            many = len(items) > 1
            parent = f"{base}/{title}" if many else base
            n_up = len({it[0] for it in items})
            for uid, fname, size, member, shown, n_run in items:
                if many:
                    # Which upload (when there are several) and which member of it
                    # (when it holds several): "Next Jumble 003", "EN_Mortadelo.nex".
                    parts = []
                    if n_up > 1:
                        parts.append(re.sub(r"\.zip$", "", _text(shown or fname), flags=re.I))
                    if n_run > 1 and member:
                        parts.append(_text(member.rsplit("/", 1)[-1]))
                    name = " - ".join(parts) or title
                else:
                    name = f"{title}  {dev}" if dev else title
                loc = f"itch:{host}/{slug}/{uid}/{quote(fname)}"
                if member:
                    loc += "?run=" + quote(member)
                tree.setdefault(parent, []).append(Entry(is_dir=False, name=name, size=size, url=loc))
            if many:
                tree.setdefault(base, []).append(Entry(is_dir=True, name=title))
            if from_db:
                games += 1
            else:
                tagged += 1
        self._save_cache()
        print(f"  nextdb: {games} games from the database, {tagged} from tags")
        if games < MIN_GAMES:
            raise SourceDown(f"nextdb: only {games} games resolved — database or itch changed?")
        for k in tree:
            # Directories first, then files, both by name; duplicate names get a suffix.
            ents = sorted(tree[k], key=lambda e: (not e.is_dir, e.name.lower()))
            seen: dict[str, int] = {}
            for e in ents:
                if e.name in seen:
                    seen[e.name] += 1
                    e.name = f"{e.name} #{seen[e.name]}"
                else:
                    seen[e.name] = 1
            tree[k] = ents
        self._tree = tree
        return tree

    def list(self, path: str) -> list[Entry]:
        tree = self._load()
        if path == "":
            return [Entry(is_dir=True, name=s) for s in SHEETS if any(k.startswith(s + "/") for k in tree)]
        if path in SHEETS:
            letters = sorted({k.split("/")[1] for k in tree if k.startswith(path + "/")},
                             key=lambda x: (x != "0-9", x))
            return [Entry(is_dir=True, name=l) for l in letters]
        return tree.get(path, [])

    def fetch(self, path: str, name: str) -> tuple[bytes, str]:
        raise FileNotFoundError("nextdb is link-only: downloads go to itch.io")
