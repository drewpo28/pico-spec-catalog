import httpx, re, sys, io, zipfile, collections, time
from urllib.parse import unquote
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/120 Safari/537.36"}
c = httpx.Client(headers=UA, follow_redirects=True, timeout=60)
P = lambda *a: (print(*a), sys.stdout.flush())
csrf_of = lambda p: (re.search(r'<meta name="csrf_token" value="([^"]+)"', p) or [None, ""])[1]
def uploads_of(p):
    out = []
    for m in re.finditer(r'<div class="upload">(.*?)</div>\s*</div>', p, re.S):
        b = m.group(1)
        uid = re.search(r'data-upload_id="(\d+)"', b)
        nm = re.search(r'class="name"[^>]*title="([^"]*)"', b) or re.search(r'class="name"[^>]*>([^<]*)<', b)
        sz = re.search(r'class="file_size"[^>]*><span>([^<]*)<', b)
        out.append((uid.group(1) if uid else None, nm.group(1) if nm else "?", sz.group(1) if sz else "?"))
    return out
# 1. csrf / cookie necessity, R2 headers
g = "https://pepperfriesgames.itch.io/tetranex"; uid = "18940718"
for label, cl, data in [("fresh client, no csrf", httpx.Client(headers=UA), {}),
                        ("fresh client, bogus csrf", httpx.Client(headers=UA), {"csrf_token": "x"})]:
    r = cl.post(f"{g}/file/{uid}", params={"source": "game_download"}, data=data)
    P(label, r.status_code, r.headers.get("content-type"), r.headers.get("content-length"), r.headers.get("transfer-encoding"), r.text[:180])
cl = httpx.Client(headers=UA)
r = cl.post(f"{g}/file/{uid}", params={"source": "game_download"}, data={"csrf_token": ""})
u = re.search(r'"url":"([^"]+)"', r.text)
if u:
    u = u.group(1).replace("\\/", "/")
    h = cl.get(u); P("R2", h.status_code, dict((k, v) for k, v in h.headers.items() if k.lower() in ("content-length","content-type","content-disposition","transfer-encoding","server")), len(h.content))
# 2. the whole Games/Demos sheets
du = c.post("https://leonisx.itch.io/zx-spectrum-next-db/download_url", data={"csrf_token": csrf_of(c.get("https://leonisx.itch.io/zx-spectrum-next-db").text)}).json()["url"]
dp = c.get(du).text; xu = uploads_of(dp)[0][0]
j = c.post("https://leonisx.itch.io/zx-spectrum-next-db/file/" + xu, params={"source": "game_download"}, data={"csrf_token": csrf_of(dp)}).json()
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(c.get(j["url"]).content))
rows = []
for sh in ("Games", "Demos", "Apps"):
    ws = wb[sh]; hdr = [x.value for x in ws[2]]
    pc = next(i for i, v in enumerate(hdr) if v == "£Digital"); tc = hdr.index("Title")
    for row in ws.iter_rows(min_row=3):
        cell = row[pc]; link = cell.hyperlink.target if cell.hyperlink else ""
        rows.append((sh, row[tc].value, str(cell.value), link))
pr = collections.Counter((sh, v.split(",")[0].strip() if v else v, "itch" if "itch.io" in l else "other") for sh, t, v, l in rows)
for k, n in sorted(pr.items(), key=lambda x: -x[1])[:40]: P("price", k, n)
cand = [r for r in rows if "itch.io" in r[3] and re.match(r"https?://[^/]+\.itch\.io/[^/?#]+/?$", r[3]) and (r[2].startswith("Free") or r[2].startswith("NYOP"))]
P("candidates (Free/NYOP, itch game page):", len(cand), collections.Counter(r[0] for r in cand))
ext = collections.Counter(); kinds = collections.Counter(); zipk = collections.Counter(); minp = collections.Counter()
zips = 0
for sh, title, price, link in cand:
    if sh == "Apps": continue
    try:
        p = c.get(link).text
    except Exception as e:
        P("ERR", link, e); continue
    m = re.search(r'"min_price":(\d+)', p); mp = int(m.group(1)) if m else None
    minp[mp] += 1
    ups = uploads_of(p)
    if ups and not ups[0][0] and mp == 0:
        try:
            d = c.post(link.rstrip("/") + "/download_url", data={"csrf_token": csrf_of(p)}).json()
            if "url" in d: ups = uploads_of(c.get(d["url"]).text)
        except Exception: pass
    names = [n for _, n, _ in ups]
    for n in names: ext[n.rsplit(".", 1)[-1].lower() if "." in n[-6:] else "(none)"] += 1
    kinds["no uploads" if not ups else ("has uid" if ups[0][0] else "no uid")] += 1
    line = f"{sh} | {title} | {price} | min={mp} | {link} | " + "; ".join(f"{n} [{s}]{'' if u else ' !uid'}" for u, n, s in ups[:6])
    P(line[:400])
    zu = [(u, n, s) for u, n, s in ups if u and n.lower().endswith(".zip") and ("kB" in s or ("MB" in s and float(s.split()[0]) < 8))]
    if zu and zips < 60 and mp == 0:
        zips += 1
        r = c.post(link.rstrip("/") + "/file/" + zu[0][0], params={"source": "game_download"}, data={"csrf_token": csrf_of(p)})
        um = re.search(r'"url":"([^"]+)"', r.text)
        if um:
            try:
                zf = zipfile.ZipFile(io.BytesIO(c.get(um.group(1).replace("\\/", "/")).content))
                nl = [n for n in zf.namelist() if not n.endswith("/")]
                top = collections.Counter(n.rsplit(".", 1)[-1].lower() for n in nl if "." in n)
                ent = [n for n in nl if n.lower().endswith((".nex", ".bas", ".snx", ".tap", ".sna", ".z80", ".bin"))]
                k = ("nex" if any(n.lower().endswith(".nex") for n in nl) else "bas" if any(n.lower().endswith(".bas") for n in nl) else "other")
                zipk[(k, len([e for e in ent if e.lower().endswith(('.nex','.bas'))]) > 1)] += 1
                P(f"   ZIP {zu[0][1]}: {len(nl)} files, dirs={len(set(n.rsplit('/',1)[0] for n in nl if '/' in n))}, exts={dict(top.most_common(8))}, entries={ent[:6]}")
            except Exception as e: P("   zip err", e)
    time.sleep(0.4)
P("min_price", minp); P("upload kinds", kinds); P("upload exts", ext.most_common(25)); P("zip kinds (entry, multiple)", zipk)
